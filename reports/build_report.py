"""Generate the editable CIT300 report and its diagrams from verified project facts.

Requires python-docx and Pillow. Word is used separately to refresh fields/render.
"""
from pathlib import Path
import math
import subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
ASSETS = OUT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
NAVY = "17324D"
BLUE = "315C7D"
PALE = "EDF3F7"
GRAY = "5A6570"
LECTURER = "[Lecturer Name]"
SUBMISSION = "[Submission Date]"
FONT = Path("C:/Windows/Fonts/times.ttf")
BOLD = Path("C:/Windows/Fonts/timesbd.ttf")


def canvas(name, width=1500, height=1100):
    image = Image.new("RGB", (width, height), "white")
    return image, ImageDraw.Draw(image), ASSETS / (name + ".png")


def font(size=34, bold=False):
    return ImageFont.truetype(str(BOLD if bold else FONT), size)


def center(draw, xy, text, size=34, bold=False, fill="#17324D"):
    draw.multiline_text(xy, text, font=font(size, bold), fill=fill, anchor="mm", align="center", spacing=7)


def box(draw, rect, title, lines=None):
    x, y, w, h = rect
    draw.rounded_rectangle((x, y, x+w, y+h), radius=12, fill="#F3F6F9", outline="#315C7D", width=3)
    if lines is None:
        center(draw, (x+w/2, y+h/2), title, 39, True)
    else:
        center(draw, (x+w/2, y+40), title, 38, True)
        draw.line((x, y+78, x+w, y+78), fill="#315C7D", width=2)
        for index, line in enumerate(lines):
            draw.text((x+22, y+90+index*40), line, font=font(32), fill="#17212B")


def arrow(draw, points, label=None, hollow=False):
    draw.line(points, fill="#315C7D", width=4)
    x, y = points[-1]
    px, py = points[-2]
    angle = math.atan2(y-py, x-px)
    wing = 19
    polygon = [(x, y), (x-wing*math.cos(angle-.5), y-wing*math.sin(angle-.5)),
               (x-wing*math.cos(angle+.5), y-wing*math.sin(angle+.5))]
    draw.polygon(polygon, fill="white" if hollow else "#315C7D", outline="#315C7D", width=3)
    if label:
        lx, ly, text = label
        bounds = draw.textbbox((lx, ly), text, font=font(30))
        draw.rectangle((bounds[0]-4, bounds[1]-3, bounds[2]+4, bounds[3]+3), fill="white")
        draw.text((lx, ly), text, font=font(30), fill="#17324D")


def save(image, path):
    image.save(path, dpi=(240, 240))
    return path


def diagrams():
    paths = {}
    im, d, path = canvas("architecture", height=1150)
    box(d, (480, 25, 540, 100), "User / console terminal")
    box(d, (370, 190, 760, 130), "ConsoleApplication + ConsoleInput\n16 options and input validation")
    arrow(d, [(750,125),(750,190)])
    box(d, (390, 390, 720, 130), "UniversitySystem\nCoordinated operations and history")
    arrow(d, [(750,320),(750,390)])
    labels = [("StudentLinkedList", "Primary records"), ("StudentHashTable", "ID lookup"), ("AVLTree", "Ordered records"),
              ("ServiceQueue", "FIFO requests"), ("ActionStack", "LIFO history"), ("CampusRouteGraph", "Locations and roads")]
    for i, (name, purpose) in enumerate(labels):
        col, row = i % 3, i // 3
        x, y = 35+col*500, 645+row*225
        box(d, (x, y, 430, 140), name+"\n"+purpose)
        if row == 0:
            arrow(d, [(750,520),(750,580),(x+215,580),(x+215,y)])
        else:
            arrow(d, [(750,520),(750,560),(x+445,560),(x+445,y+70),(x+430,y+70)])
    center(d, (750,1070), "All six structures are owned directly by UniversitySystem.\nSearchAlgorithms and SortingAlgorithms operate on record snapshots.", 32)
    paths["architecture"] = save(im, path)

    im, d, path = canvas("uml_students", height=1040)
    box(d, (35,70,650,365), "Student", ["- studentId, name, age, gender", "- degreeProgramme, email", "- contactNumber, address, gpa", "+ validation methods / getters", "{immutable record}"])
    box(d, (870,70,585,275), "StudentNode", ["~ student : Student", "~ next : StudentNode", "{package-private node}"])
    box(d, (870,570,585,340), "StudentLinkedList", ["- head, tail : StudentNode", "- size : int", "+ add(), update(), delete()", "+ search(), toArray()", "+ size(), isEmpty()"])
    box(d, (35,570,650,340), "StudentHashTable", ["- buckets : Entry[]; size : int", "{Entry: Student, next}", "+ insert(), search(), delete()", "+ update(), describeBuckets()", "- hash(), resize()"])
    arrow(d, [(870,185),(685,185)], (705,137,"record 1"))
    arrow(d, [(1160,570),(1160,345)], (1180,440,"links 0..*"))
    arrow(d, [(355,570),(355,435)], (370,475,"indexes 0..*"))
    center(d, (750,985), "Arrow = navigable reference; ~ = package access; + = public; - = private.\nStudentNode.next links to another StudentNode or null.", 31)
    paths["uml_students"] = save(im, path)

    im, d, path = canvas("uml_trees", height=1030)
    box(d, (45,80,650,350), "StudentBST", ["# root : TreeNode; - size : int", "+ insert(), search(), update()", "+ delete(), inorder(), height()", "# restore(node), refreshHeight()", "{key: normalized Student ID}"])
    box(d, (905,80,540,350), "TreeNode", ["~ student : Student", "~ left, right : TreeNode", "~ height : int", "{null children at leaves}"])
    box(d, (45,660,650,265), "AVLTree", ["# restore(node) : TreeNode", "- rotateLeft(), rotateRight()", "+ isValidAvl()", "{extends StudentBST}"])
    arrow(d, [(695,245),(905,245)], (710,188,"root 0..1"))
    arrow(d, [(370,660),(370,430)], (400,515,"inherits"), hollow=True)
    center(d, (1145,660), "TreeNode.left and .right\nreference TreeNode (0..1).\n\nEach node references\none immutable Student.\n\nAVL balances after\ninsertion and deletion.", 36)
    center(d, (750,985), "Hollow triangle = inheritance; # = protected; other symbols as in the student diagram.", 31)
    paths["uml_trees"] = save(im, path)

    im, d, path = canvas("uml_queue_stack", height=1180)
    for x, title, fields, node, value, valuefields in [
        (35,"ServiceQueue",["- front, rear : QueueNode","- size : int","+ enqueue(), dequeue(), peek()","+ toArray(), hasStudent()"],"QueueNode","ServiceRequest",["- requestId, studentId","- requestType, description","+ getters / toString()"]),
        (815,"ActionStack",["- top : StackNode","- size : int","+ push(), pop(), peek()","+ toArray(), size(), isEmpty()"],"StackNode","Action",["- description : String","- timestamp : LocalDateTime","+ getters / toString()"])
    ]:
        box(d,(x,30,650,310),title,fields)
        box(d,(x,500,650,210),node,["~ "+("request : ServiceRequest" if node=="QueueNode" else "action : Action"),"~ next : "+node])
        box(d,(x,865,650,245),value,valuefields)
        arrow(d,[(x+325,340),(x+325,500)],(x+345,390,"owns links"))
        arrow(d,[(x+325,710),(x+325,865)],(x+345,765,"record 1"))
    paths["uml_queue_stack"] = save(im,path)

    im,d,path=canvas("uml_graph",height=1040)
    box(d,(35,70,670,380),"CampusRouteGraph",["- locations : Location[]","- size, roadCount : int","+ addLocation(), removeLocation()","+ addRoad(), removeRoad()","+ bfs(), dfs(), describeConnections()","- append(), unlink(), findEdge()"])
    box(d,(925,70,530,335),"Location",["- name : String","~ index : int","~ firstEdge, lastEdge : Edge","+ getName()"])
    box(d,(925,650,530,260),"Edge",["- destination : Location","- distance : double","~ next : Edge","+ getters"])
    arrow(d,[(705,230),(925,230)],(725,172,"0..*"))
    arrow(d,[(1100,405),(1100,650)],(930,500,"adjacency 0..*"))
    arrow(d,[(1350,650),(1350,405)],(1370,500,"1"))
    center(d,(370,720),"Each road is represented\nby two directed Edge objects.\n\nEdge.next links the adjacency list.\nLocation indexes change when\nthe vertex array is compacted.",34)
    center(d,(750,980),"Association arrows follow actual object references; private links are graph-owned.",31)
    paths["uml_graph"] = save(im,path)

    im,d,path=canvas("system_flow",height=1470)
    items=[("User input",35),("Main menu (1-16)",225),("Validate selection and fields",415),("Dispatch to selected system module",605),("Operate on custom data structures",795),("Display result / record successful action",985)]
    for i,(label,y) in enumerate(items):
        box(d,(260,y,920,110),label)
        if i: arrow(d,[(720,items[i-1][1]+110),(720,y)])
    arrow(d,[(1180,1040),(1350,1040),(1350,280),(1180,280)],(1190,700,"repeat"))
    box(d,(380,1270,680,110),"Option 16 or closed input: exit")
    arrow(d,[(260,280),(125,280),(125,1325),(380,1325)],(135,1110,"exit path"))
    center(d,(720,1190),"Invalid data is explained and re-entered; no partial student update is committed.",29)
    paths["flow"] = save(im,path)

    im,d,path=canvas("campus_graph",width=1600,height=960)
    pos={"Library":(190,450),"Main Gate":(470,140),"Engineering\nFaculty":(955,140),"Lecture Hall":(1295,450),"Computer\nLaboratory":(745,450),"Cafeteria":(745,785),"Hostel":(1310,785)}
    edges=[("Main Gate","Library",150), ("Main Gate","Engineering\nFaculty",300), ("Library","Computer\nLaboratory",100), ("Library","Cafeteria",120), ("Engineering\nFaculty","Lecture Hall",80), ("Computer\nLaboratory","Lecture Hall",90), ("Lecture Hall","Cafeteria",140), ("Cafeteria","Hostel",250)]
    for a,b,weight in edges:
        x,y=pos[a]; xx,yy=pos[b]
        d.line((x,y,xx,yy),fill="#315C7D",width=5)
        mid=((x+xx)/2,(y+yy)/2-22)
        label=str(weight)+" m"
        bb=d.textbbox(mid,label,font=font(35),anchor="mm")
        d.rectangle((bb[0]-8,bb[1]-5,bb[2]+8,bb[3]+5),fill="white")
        center(d,mid,label,35)
    for label,(x,y) in pos.items():
        box(d,(x-150,y-58,300,116),label)
    center(d,(800,915),"Undirected illustrative campus graph: 7 vertices, 8 roads; not a geographical map.",32)
    paths["campus"] = save(im,path)
    return paths


document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.85)
section.left_margin = section.right_margin = Inches(.95)
section.header_distance = section.footer_distance = Inches(.4)
for name in ["Normal","Body Text","Caption","Title","Subtitle","Heading 1","Heading 2","Heading 3","List Bullet","List Number"]:
    style=document.styles[name]
    style.font.name="Times New Roman"
    style.font.size=Pt(12)
    style.paragraph_format.line_spacing=1.5
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.widow_control=True
    style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [("Heading 1",16),("Heading 2",14),("Heading 3",14)]:
    st=document.styles[name]
    st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY)
    st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
document.styles["Heading 1"].paragraph_format.page_break_before=True
document.styles["Caption"].font.italic=True
document.styles["Caption"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
document.styles["Caption"].paragraph_format.space_after=Pt(10)


def shade(cell, fill):
    element=OxmlElement("w:shd"); element.set(qn("w:fill"),fill); cell._tc.get_or_add_tcPr().append(element)


def field(paragraph, instruction, display="1"):
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    code=OxmlElement("w:instrText"); code.set(qn("xml:space"),"preserve"); code.text=" "+instruction+" "
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    paragraph.add_run()._r.append(begin); paragraph.add_run()._r.append(code); paragraph.add_run()._r.append(sep)
    paragraph.add_run(display); paragraph.add_run()._r.append(end)


def p(text, bold=False):
    paragraph=document.add_paragraph()
    paragraph.add_run(text).bold=bold
    return paragraph


def bullets(items):
    for text in items: document.add_paragraph(text,style="List Bullet")


def h(text,level=2):
    return document.add_heading(text,level)


def chapter(number,title):
    return h("CHAPTER "+str(number)+": "+title.upper(),1)


def caption(label,text):
    paragraph=document.add_paragraph(style="Caption")
    paragraph.add_run(label+" "); field(paragraph,"SEQ "+label+" \\* ARABIC")
    paragraph.add_run(". "+text)
    return paragraph


def table(headers, rows, title=None, widths=None):
    if title:
        cap=caption("Table",title); cap.paragraph_format.keep_with_next=True
    t=document.add_table(rows=1,cols=len(headers)); t.style="Table Grid"
    t.autofit=False
    if widths:
        for column,width in zip(t.columns,widths): column.width=Inches(width)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=text; shade(cell,NAVY)
        for run in cell.paragraphs[0].runs:
            run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    repeat=OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for idx, row in enumerate(rows):
        cells=t.add_row().cells
        for cell,text in zip(cells,row):
            cell.text=str(text)
            if idx%2==0: shade(cell,"F4F7FA")
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after=Pt(3)
                paragraph.paragraph_format.space_before=Pt(3)
                paragraph.paragraph_format.line_spacing=1.5
                for run in paragraph.runs: run.font.name="Times New Roman"; run.font.size=Pt(12)
    document.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def figure(path,title,width=6.1):
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next=True
    paragraph.add_run().add_picture(str(path),width=Inches(width))
    caption("Figure",title)


def placeholder(title,instruction,height=.65,numbered=True):
    t=document.add_table(rows=1,cols=1)
    t.style="Table Grid"; shade(t.cell(0,0),PALE)
    t.rows[0].height=Inches(height)
    paragraph=t.cell(0,0).paragraphs[0]
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run("[ SCREENSHOT PLACEHOLDER ]\n"); run.bold=True; run.font.color.rgb=RGBColor.from_string(BLUE)
    paragraph.add_run(instruction)
    for run in paragraph.runs: run.font.name="Times New Roman"; run.font.size=Pt(12)
    paragraph.paragraph_format.keep_with_next=numbered
    if numbered: caption("Figure",title+" (screenshot placeholder)")


def page(): document.add_page_break()


def footer(sec,roman=False):
    sec.header.is_linked_to_previous=False; sec.footer.is_linked_to_previous=False
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=header.add_run("CIT300  |  Graded Practical Assignment 1"); run.font.name="Times New Roman"; run.font.size=Pt(10); run.font.color.rgb=RGBColor.from_string(GRAY)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(fp,"PAGE")
    pn=OxmlElement("w:pgNumType"); pn.set(qn("w:start"),"1"); pn.set(qn("w:fmt"),"lowerRoman" if roman else "decimal")
    sec._sectPr.append(pn)


def fronttitle(text):
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
    run=paragraph.add_run(text); run.bold=True; run.font.size=Pt(16); run.font.color.rgb=RGBColor.from_string(NAVY)


paths=diagrams()
document.core_properties.title="University Student Record and Campus Route Management System"
document.core_properties.subject="CIT300 - Data Structures and Algorithms | Graded Practical Assignment 1"
document.core_properties.author="J. Abdullah; SMF. Asra; MSM. Dilsath; MTF. Nifra"
document.core_properties.keywords="CIT300, Java, linked list, queue, stack, AVL, hashing, graph, BFS, DFS"

# Cover.
for text,size,bold in [
    ("SRI LANKA TECHNOLOGICAL CAMPUS (SLTC)",16,True),
    ("CIT300 - DATA STRUCTURES AND ALGORITHMS",14,True),
    ("Graded Practical Assignment 1",14,True),
    ("UNIVERSITY STUDENT RECORD AND\nCAMPUS ROUTE MANAGEMENT SYSTEM",16,True),
    ("Java Console-Based Application",12,False)
]:
    para=document.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(16)
    run=para.add_run(text); run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(NAVY)
para=p("Prepared By"); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.runs[0].bold=True
members=[("J. Abdullah","23DA2-0575","Student Linked List"),("SMF. Asra","23DA2-0826","Queue and Stack"),("MSM. Dilsath","23DA2-0576","Hash Table and BST/AVL"),("MTF. Nifra","23DA2-0729","Campus Graph Routes")]
table(["Group Member","Student ID"],[(a,b) for a,b,c in members],widths=[3.6,2.6])
for text in ["Submitted To: "+LECTURER,"Submission Date: "+SUBMISSION]:
    para=p(text); para.alignment=WD_ALIGN_PARAGRAPH.CENTER

front=document.add_section(WD_SECTION_START.NEW_PAGE); footer(front,True)
fronttitle("TABLE OF CONTENTS")
field(p(""),'TOC \\o "1-2" \\h \\z \\u',"Contents will be generated by Microsoft Word.")
page(); fronttitle("LIST OF FIGURES")
field(p(""),'TOC \\h \\z \\c "Figure"',"Figure list will be generated by Microsoft Word.")
page(); fronttitle("LIST OF TABLES")
field(p(""),'TOC \\h \\z \\c "Table"',"Table list will be generated by Microsoft Word.")

body=document.add_section(WD_SECTION_START.NEW_PAGE); footer(body)
first=chapter(1,"Introduction"); first.paragraph_format.page_break_before=False
h("1.1 Background")
p("Universities maintain information about students, programmes, academic results and administrative requests. When these records are handled through unrelated documents or repeated manual entry, finding an individual record and keeping revised information consistent becomes difficult. A structured student management system provides a common way to create, locate, update and display the information needed for routine administration.")
p("A campus also contains interconnected locations such as a library, laboratories, lecture halls and accommodation. Describing these locations as an ordinary list does not explain which places are connected. A graph captures both locations and roads, allowing traversal algorithms to explore the routes available from a selected starting point.")
h("1.2 Role of data structures")
p("This project treats student administration and campus connectivity as practical data-structure problems. A linked list keeps records in insertion order; a queue preserves the order of service requests; a stack shows the most recent activity first. Hashing supports student-ID lookup, an AVL tree maintains ordered records, and a graph supports campus traversal. These choices make the consequences of each structure visible through one console interface.")
h("1.3 Scope of the implemented application")
p("The developed Java application offers the required 16-option main menu. Every main data structure is implemented using custom nodes or arrays, without Java Collection Framework storage classes. The system runs with a Java 17 language target and separates console input, presentation and coordinated data operations into distinct classes.")
p("Records are stored in memory for one execution. A normal launch starts empty; the --demo option loads four sample students and seven campus locations connected by eight illustrative roads. Service processing is simulated, and BFS/DFS report reachability and traversal order rather than a distance-optimal journey. These boundaries keep the work focused on the assignment's DSA objectives.")

chapter(2,"Project Objectives")
p("The overall objective is to implement a coherent university management demonstration in which each data structure has a clear responsibility and all operations are accessible through a validated Java console interface.")
table(["Objective","Implementation and evidence"],[
    ("Manage student records","Create, update, delete and display nine-field Student records through a custom singly linked list."),
    ("Apply data structures practically","Implement custom list, queue, stack, hash table, BST/AVL and graph classes."),
    ("Support fast ID lookup","Use a custom hash function with separate chaining and automatic resizing."),
    ("Organize student information","Use AVL inorder traversal for Student ID order; use merge sort for name and GPA views."),
    ("Manage administrative requests","Enqueue validated requests and process the oldest request first."),
    ("Represent campus routes","Manage graph vertices and weighted bidirectional roads, including removal of incident roads."),
    ("Demonstrate traversal","Run BFS with an array queue and DFS with an array stack from a chosen location."),
    ("Verify reliable behavior","Check invalid inputs, empty structures and cross-index consistency using automated tests.")
],"Project objectives and implementation evidence",[2.1,4.1])
h("2.1 Success criteria")
p("Success requires compilation with the Java 17 target, access to all 16 menu choices, correct ordering behavior in each structure, and consistent student contents across the linked list, hash table and AVL tree. Tests must also demonstrate that invalid input does not introduce partial student records or break graph connections.")

chapter(3,"System Overview")
h("3.1 Architecture")
p("Main creates UniversitySystem and ConsoleApplication. ConsoleInput reads whole lines, validates the input and returns typed values. ConsoleApplication dispatches menu choices, while UniversitySystem owns the structures and coordinates successful changes. This separation keeps the console menu independent from the linked-list, tree, hashing and graph algorithms.")
figure(paths["architecture"],"Application architecture and module responsibilities",6.05)
h("3.2 Module responsibilities")
table(["Module","Classes and role"],[
    ("Student records","Student and StudentLinkedList hold validated immutable records and their insertion order."),
    ("Service requests","ServiceRequest and ServiceQueue implement first-in, first-out request handling."),
    ("Recent history","Action and ActionStack store timestamped operations in last-in, first-out order."),
    ("Student organization","StudentBST defines ordered operations; AVLTree supplies balancing for the running application."),
    ("Hash lookup","StudentHashTable maintains a separate student-ID index with collision chains."),
    ("Campus routes","CampusRouteGraph owns Location vertices, Edge adjacency lists and BFS/DFS operations.")
],"System modules",[1.8,4.4])
h("3.3 Coordinated record lifecycle")
p("A successful addition validates a Student and inserts the same immutable object into the list, hash table and AVL tree. Updating constructs a complete replacement and replaces the corresponding record in all three structures. Deletion removes the ID from all indexes after confirming that no waiting request still refers to that student. Successful changes are recorded in ActionStack.")
p("Returned arrays are snapshots. Sorting a snapshot does not reorder the primary linked list. This allows the same student set to be presented in insertion order, ID order, alphabetical order or GPA order without compromising the underlying storage sequence.")

chapter(4,"Technology Used")
table(["Technology","Use in this project"],[
    ("Java / JDK","Java 17 language and bytecode target. Local verification used JDK 23.0.2 with --release 17."),
    ("Visual Studio Code","Source editing, project navigation and supplied Java launch configurations."),
    ("Git and GitHub","Local commits and branches are present. GitHub workflow and pull-request template are prepared; remote evidence is pending."),
    ("PowerShell / shell scripts","Dependency-free compilation, application launch and automated testing."),
    ("Standard JDK APIs","BufferedReader and streams for console I/O; java.time for timestamps; Locale for consistent formatting [1].")
],"Technology stack and actual usage",[1.7,4.5])
h("4.1 Object-oriented programming")
p("Encapsulation keeps data-structure state behind private fields and public operations. Student, ServiceRequest and Action represent domain values. UniversitySystem uses composition to own the application structures. Inheritance is demonstrated by AVLTree extending StudentBST and overriding the balancing hook while reusing the base search, update and traversal behavior.")
h("4.2 Custom implementation constraint")
p("The project does not use ArrayList, LinkedList, HashMap, Stack, Queue, TreeMap, collection-based sorting, or graph libraries. Nodes, bucket arrays, traversal arrays and merge-sort buffers are manually managed. java.util.Locale is used for text and numeric formatting only; it is not a collection. The project-specific algorithms are explained from the source rather than attributed to library implementations.")
h("4.3 Execution")
p("From the project folder, run: powershell -NoProfile -ExecutionPolicy Bypass -File scripts/build.ps1 run -Demo. On a POSIX-compatible shell, run: sh scripts/build.sh run --demo. After compilation, java -cp out Main --demo launches the demonstration directly. Omit --demo to start with empty structures.")

chapter(5,"System Requirements Analysis")
h("5.1 Functional requirements")
table(["ID","Requirement","Implemented behavior"],[
    ("FR01","Student creation","Validate nine fields, reject duplicate IDs and insert in list/hash/AVL."),
    ("FR02","Student update","Locate ID, preserve unchanged fields and replace all indexed references."),
    ("FR03","Student deletion","Remove from all three structures; block deletion while requests remain."),
    ("FR04","Display and search","Show complete records; support hash, linear and binary lookup."),
    ("FR05","Sorting and ordering","Show AVL ID order, alphabetical names and descending GPA ranking."),
    ("FR06","Service requests","Add, display and process requests using FIFO semantics."),
    ("FR07","Action history","Push successful actions, display newest first, and allow peek/pop."),
    ("FR08","Campus locations","Add unique location names; remove vertices and incident roads."),
    ("FR09","Campus roads","Add/remove bidirectional positive-distance connections and display adjacency."),
    ("FR10","Campus traversal","Select BFS or DFS and a starting location; report reachable vertices."),
    ("FR11","Safe console control","Reject invalid menu choices, explain missing data and exit on EOF.")
],"Functional requirements",[.65,1.65,3.9])
h("5.2 Non-functional requirements")
bullets(["Maintainability: separate packages for domain records, structures, algorithms and console coordination.","Portability: compile for Java 17 without Maven, Gradle or third-party runtime libraries.","Usability: retain the required menu labels and provide readable validation messages.","Correctness: preserve list/hash/tree agreement and graph adjacency invariants after mutations.","Demonstrability: provide deterministic traversal order, sample data, tests and a documented startup procedure."])
h("5.3 Data requirements")
p("Student IDs are strings, preserving leading zeros in 0001–0004. The remaining fields are name, age, gender, programme, email, contact number, address and GPA. Request records contain request ID, student ID, type and description. Graph roads contain two existing endpoints and a positive distance in metres.")

chapter(6,"System Design")
h("6.1 UML class diagrams")
p("The diagrams show selected fields and operations from the implemented classes. They are divided by responsibility to remain readable. Navigable references are shown by arrows, and the hollow triangle identifies inheritance. Node-to-node links may be null at the end of a list or at a tree leaf.")
figure(paths["uml_students"],"UML: student records, linked list and hash index",6.15)
page(); h("6.1.1 Tree classes")
p("StudentBST owns a root reference and delegates rebalancing to restore(). AVLTree overrides that operation. TreeNode stores the Student reference, child links and cached height. Student ID is the ordering key throughout the tree hierarchy.")
figure(paths["uml_trees"],"UML: BST and AVL inheritance",6.15)
page(); h("6.1.2 Queue and stack classes")
p("QueueNode and StackNode are internal links. Their containing structures expose operations without exposing mutable node references. ServiceRequest and Action encapsulate the payloads carried by those links.")
figure(paths["uml_queue_stack"],"UML: service queue and action stack",6.1)
page(); h("6.1.3 Graph classes")
p("Location objects keep their name, current array index and adjacency-list endpoints. An Edge identifies one destination and distance. CampusRouteGraph maintains the vertex array, creates a pair of edges for every road and removes both directions together.")
figure(paths["uml_graph"],"UML: campus graph classes and references",6.1)
page(); h("6.2 System flow diagram")
p("The menu drives a repeated input–validate–execute–display cycle. Modules perform the requested operation and return results to the console. Option 16 exits normally; an end-of-input condition is handled by Main. The coordinator only applies a new or updated Student after all fields have been collected successfully.")
figure(paths["flow"],"System flow from user input to custom structures",5.35)
h("6.3 Design decisions")
p("The linked list is the primary record sequence; the hash table and AVL tree are secondary indexes. The ID remains fixed during editing so the key is stable across all structures. Graph names are case-insensitive, and the graph is undirected because the current model treats each campus road as usable in both directions. The history stack records actions but does not implement undo.")

chapter(7,"Data Structure Implementation")
h("7.1 Linked list implementation")
p("StudentLinkedList stores a head, tail and size. Each StudentNode contains a Student and a next reference. The head identifies the first record, the tail supports constant-time link attachment, and next defines insertion-order traversal. The list itself is responsible for maintaining these relationships when records change.")
bullets(["Insertion: reject null or an existing ID, create a node, attach it after tail, and update head when the list was empty. Increase size once.","Searching: normalize the supplied ID and move from head until a matching student is found or the end is reached.","Updating: locate the node with the same ID and replace its Student reference. Node position is preserved.","Deletion: retain previous/current references, bypass the matching node, move head for first-node deletion and move tail for last-node deletion.","Display: toArray() traverses links and returns the Student references in insertion order for the console to print."])
p("Search, update and deletion take O(n) time. Although attachment at the tail is O(1), add() performs a duplicate search and therefore has O(n) overall time. Empty-list deletion returns null. Deleting the sole node resets both endpoints, allowing the list to be reused correctly.")
placeholder("Linked-list source code","Insert VS Code screenshot of StudentNode and StudentLinkedList.add/delete.",1.0)

page(); h("7.2 Queue implementation")
p("ServiceQueue uses linked QueueNode objects with front and rear references. It models FIFO: a request arriving earlier is processed earlier. ServiceRequest holds a generated request ID, a validated student ID, request type and description.")
bullets(["Enqueue: create a node, link it after rear and advance rear. Set front as well when inserting into an empty queue.","Dequeue: return the front request and advance front. If the queue becomes empty, reset rear to null.","Peek: return the next request without removing it.","Display: traverse from front to rear and show the waiting requests in processing order."])
p("Enqueue, dequeue and peek are O(1). Display and the hasStudent() check take O(q). UniversitySystem checks that the student exists before adding a request and allocates sequential IDs such as REQ-1. It prevents deletion of a student whose requests are still pending; processing the requests first avoids orphan references.")
h("7.3 Stack implementation")
p("ActionStack holds a top reference to linked StackNode objects. Push creates a node whose next link is the old top, then makes the new node top. Pop returns the top Action and moves top to the next node. Peek leaves the structure unchanged. All three operations take O(1); display traverses O(h) entries from newest to oldest.")
p("Each Action contains a description and LocalDateTime timestamp. Student changes, searches, requests and campus changes are recorded. Empty pop/peek operations return null and produce an appropriate message. Popping an entry removes that history record only; it neither restores deleted records nor reverses roads.")

page(); h("7.4 BST / AVL tree implementation")
p("StudentBST organizes TreeNode objects by normalized Student ID. Each node has a Student, left and right child references, and a cached height. Search follows the comparison with the current key, descending left for a smaller ID and right for a larger ID. Equal IDs represent the same record and duplicate insertion is rejected.")
p("Insertion recursively locates a null child position and places a new node there. Deletion handles three cases: a leaf becomes null; a one-child node is replaced by its child; and a two-child node receives the smallest record from its right subtree before that successor is removed. Inorder traversal visits left subtree, node and right subtree, producing ascending Student ID order.")
p("The running application uses AVLTree. After insertion or deletion, restore() updates height and calculates left-height minus right-height. A balance outside -1 to +1 triggers a single or double rotation. LL and RR imbalances use one rotation; LR and RL cases first rotate the child. Search and updates retain O(log n) tree height, while inorder traversal visits all n nodes.")
p("The unbalanced BST remains available as a comparison implementation and is tested separately. The AVL validator independently checks key ordering, stored heights and subtree balance. GPA ranking uses merge sort rather than changing the tree key; equal GPAs receive the same competition rank.")
table(["Operation","BST","AVL used by application"],[("Search / insert / delete","O(height); worst O(n)","O(log n)"),("Inorder display","O(n)","O(n)"),("Balancing","None","Single/double rotations")],"Tree behavior and complexity",[2.0,2.0,2.2])

page(); h("7.5 Hash table implementation")
p("StudentHashTable owns an array of bucket references. Each bucket is the head of a custom Entry chain, and each Entry stores a Student and next reference. A polynomial hash processes the characters of the normalized ID using hash = (31 × hash + character) mod capacity. A long intermediate value is used during multiplication.")
p("Insertion first rejects a duplicate ID. If the resulting load would exceed 0.75, the table expands from capacity c to 2c + 1 and relinks entries into newly calculated buckets. The new student is then inserted at the head of its bucket chain. Searching hashes the ID and compares entries only within the selected chain.")
p("Deletion tracks the previous entry while walking a bucket. It either replaces the bucket head or bypasses the matching entry. Updating replaces the matching Student reference without changing the key. Separate chaining permits several different IDs to share the same bucket without overwriting one another.")
p("Expected lookup, insertion and deletion are O(1) under a suitable key distribution and bounded ID length; worst-case collisions give O(n). A resize takes O(n), so the expected insertion bound is amortized. Tests intentionally place AA, AR, RA and RR in the same capacity-17 bucket, then exercise deletion at the head, middle and tail.")
h("7.5.1 Searching and sorting support")
p("Linear search visits the record snapshot sequentially. Binary search first creates a copy sorted by ID or name, then narrows the candidate interval. Name searching returns all exact matches without case sensitivity. A binary search phase is O(log n), but preparing its sorted copy makes the complete public operation O(n log n).")
p("Merge sort splits a snapshot into halves, recursively sorts the halves and merges them through an auxiliary array. It supports name ascending, GPA descending and ID ascending. Tied names or GPAs use ID as a deterministic secondary key. The algorithm takes O(n log n) time and O(n) auxiliary space, and the original linked list is unchanged.")

page(); h("7.6 Campus route graph implementation")
p("CampusRouteGraph stores Location vertices in an expandable array and roads as linked Edge adjacency lists. A road is represented by two directed edge objects with the same positive distance. This gives an undirected weighted model while keeping adjacency traversal straightforward.")
figure(paths["campus"],"Implemented demonstration campus graph and distances",6.15)
bullets(["Add location: trim and validate the name, reject a case-insensitive duplicate and expand the vertex array if needed.","Remove location: unlink all incoming references, discard its own adjacency links, compact the vertex array and update remaining indexes.","Add connection: verify two distinct endpoints and a finite positive distance; reject an existing road and append both edge directions.","Remove connection: unlink the edge from each endpoint and decrement roadCount once.","Display connections: list every vertex and its neighbors with distances, including isolated locations."])
page(); h("7.6.1 Breadth-first search")
p("BFS allocates a boolean visited array and an integer array queue. It marks the start before enqueueing it, repeatedly removes the oldest queued vertex and enqueues each unvisited neighbor. Marking on discovery prevents duplicate queue entries in a cyclic graph. The traversal follows breadth levels from the starting location.")
p("For the demo graph starting at Main Gate, the order is: Main Gate → Library → Engineering Faculty → Computer Laboratory → Cafeteria → Lecture Hall → Hostel. The order depends on adjacency insertion order, which the implementation preserves.")
h("7.6.2 Depth-first search")
p("DFS uses a manual stack of Edge cursors to simulate recursive exploration. The top cursor advances through a vertex's neighbors; when an unvisited neighbor is found, its adjacency cursor is pushed. When a cursor becomes null, the algorithm backtracks by popping. This design avoids recursion in graph traversal.")
p("For the same demo start, the order is: Main Gate → Library → Computer Laboratory → Lecture Hall → Engineering Faculty → Cafeteria → Hostel. Each vertex appears once because a visited array guards discoveries. Both traversals use O(V + E) time and O(V) working memory for the graph model.")
h("7.6.3 Traversal limits and overall costs")
p("Only the connected component reachable from the selected start is traversed. A disconnected location remains visible in the adjacency display but is absent from that traversal. The console reports visited locations against the total vertex count. Distances are displayed but do not alter BFS/DFS order; weighted shortest-path computation would require an additional algorithm.")
table(["Operation","Time cost"],[("List-backed student add/update","O(n), including list traversal"),("Student deletion with pending-request check","O(n + q)"),("Queue/stack core operations","O(1)"),("Graph location lookup/add","O(V), including name search"),("Graph vertex removal or adjacency display","O(V + E)"),("Graph BFS / DFS","O(V + E), O(V) working space")],"Selected application-level complexity",[4.25,1.95])

chapter(8,"System Menu and Functionality")
p("The main menu preserves all 16 required labels. The following entries explain the operation and identify the screenshot to capture during the final demonstration. Screenshot boxes are intentionally editable placeholders, not fabricated interface evidence. Option 4, 5 and 7 expose related operations through submenus.")
menus=[
    ("Add Student Record","Enter an unused ID and all remaining fields. The application validates the complete record, inserts it into the linked list, hash table and AVL tree, then records an action. A duplicate ID is rejected before collecting the other fields.","Capture a valid new record and the message confirming insertion in all three indexes."),
    ("Update Student Record","Enter an existing ID. Current values are displayed, and Enter keeps an existing value. The ID stays fixed. After all revised fields are valid, the new immutable record replaces the old reference in every index.","Capture a changed GPA or programme followed by the successful update message."),
    ("Delete Student Record","Enter the ID to remove. The system rejects a missing record or a student with waiting service requests. Successful deletion updates all three indexes and pushes a history entry.","Capture deletion and a subsequent lookup showing that the ID is absent."),
    ("Display All Records using Linked List","The console displays all nine fields in insertion order. Its record-tools submenu provides linear and binary search by ID/name, alphabetical sorting, GPA ranking and hash bucket display. Use 0 to return.","Capture the four sample records or a complete single-record view and record-tools menu."),
    ("Add Service Request to Queue","Choose submenu 1, enter an existing student ID, request type and description. A generated request ID is enqueued. Submenu 2 displays all waiting requests; submenu 0 returns to the main menu.","Capture two waiting requests in the order they were entered."),
    ("Process Next Service Request","The front request is removed and shown as processed. The simulated processing follows FIFO. If the queue is empty, a message explains that no requests are waiting.","Capture the oldest request being processed and the remaining queue."),
    ("Display Recent Actions using Stack","History appears newest first. The submenu supports peek, pop and redisplay. Pop removes only the top history entry; it does not undo the original data change.","Capture recent actions and the result of peek or pop."),
    ("Display Students using BST/AVL","The AVL inorder traversal prints student ID, name, programme and GPA in ascending ID order. No records are inserted or reordered during this read-only display.","Capture IDs 0001, 0002, 0003 and 0004 in ascending order."),
    ("Search Student using Hashing","Enter a student ID. The hash table selects a bucket, searches its collision chain and returns the matching complete record. An unsuccessful lookup displays a clear message; the search is logged.","Capture lookup of 0002 and its matching name, Asra."),
    ("Add Campus Location","Enter a nonblank unique name. A new graph vertex is added and an action is pushed. Matching is case insensitive, so Library and library identify the same location.","Capture adding a new location such as Sports Centre."),
    ("Remove Campus Location","Enter an existing location name. The graph removes the vertex and every incident road, then updates vertex indexes used by traversal.","Capture removal and an adjacency display without dangling roads."),
    ("Add Campus Connection/Road","Enter two existing, different endpoint names and a finite positive distance in metres. Two adjacency entries represent one bidirectional road. Reverse duplicates and self-loops are rejected.","Capture a valid road addition with its distance."),
    ("Remove Campus Connection/Road","Enter the two endpoint names. The graph removes both adjacency directions and records one road removal. A missing connection is reported without changing other roads.","Capture removal followed by the updated adjacency list."),
    ("Display Campus Connections","The adjacency display lists each location, neighboring locations and distances. Isolated vertices are shown with no roads; the final line gives vertex and road counts.","Capture the seven-location, eight-road demo map."),
    ("Traverse Campus Locations using BFS or DFS","Choose 1 for BFS or 2 for DFS and enter a starting location. The application prints a deterministic traversal order and the reached-versus-total location count.","Capture BFS and DFS from Main Gate, including their different visit orders."),
    ("Exit","Selecting 16 exits the loop and prints Goodbye! Closing standard input is also handled cleanly by Main. Session data is discarded when the process ends.","Capture option 16 and the final Goodbye! message.")
]
for index,(title,_,__) in enumerate(menus,1):
    para=p(str(index)+". "+title)
    para.paragraph_format.space_after=Pt(0)
for index,(title,explanation,instruction) in enumerate(menus,1):
    if index%2==1: page()
    h("8."+str(index)+" "+title)
    p(explanation)
    placeholder("Menu option "+str(index)+": "+title,instruction,.72)

chapter(9,"Input Validation and Error Handling")
p("Validation is applied both at the console and inside the model/structure APIs. ConsoleInput reads complete lines and catches invalid numeric or domain values so that the field can be entered again. UniversitySystem performs existence and cross-record checks before modifying the structures.")
table(["Condition","Handling and result"],[
    ("Duplicate student ID","Trim and uppercase the ID, then reject a match before insertion."),
    ("Invalid student information","Reject empty required fields; require age 1–120, finite GPA 0–4, basic email syntax and a 7–15 digit contact number."),
    ("Invalid menu selection","Reject nonnumeric/out-of-range input and request a permitted choice."),
    ("Missing student record","Show a clear missing-record message; do not change list, hash or tree."),
    ("Invalid campus connection","Reject missing endpoints, self-loops, duplicates and nonpositive/nonfinite distances."),
    ("Empty structures","Return null or an empty array as appropriate; print readable empty-state messages."),
    ("Interrupted input","Handle EOF in Main. An unfinished Student is never committed to the indexes."),
    ("Pending service request","Prevent deletion of its student until queued requests have been processed.")
],"Input validation and error-handling rules",[2.0,4.2])
h("9.1 Consistency under failure")
p("Immutable Student construction means an invalid field cannot partially modify an existing record. Routine validation errors are represented by IllegalArgumentException and displayed by the console. IOException and EOFException are handled at the application boundary. Unexpected implementation faults are not hidden by a blanket catch, allowing tests to expose genuine defects.")
p("Email checks validate a basic shape only and do not establish that a mailbox exists. The console application does not provide authentication, durable storage or concurrent-user access. These capabilities are outside the current assignment scope.")

chapter(10,"GitHub Collaboration and Version Control")
h("10.1 Repository and current evidence")
p("Repository name: University-Student-Record-Campus-Route-Management-System. The local repository has an initial implementation commit, 4128be8, and the branch names listed below. Inspection during report preparation found no configured remote. A README URL, workflow file or planned contribution is not evidence that a repository has been published or that a pull request has been merged.")
p("The report therefore distinguishes the requested GitHub collaboration workflow from verified local state. Remote repository, contributor and pull-request screenshots remain placeholders until the group supplies the actual evidence.")
table(["Requested workflow branch","Existing local equivalent"],[("main","main"),("student-management","student-record-module"),("queue-stack-management","queue-stack-module"),("hashing-tree-management","hash-tree-module"),("campus-route-management","campus-route-module")],"Requested branch names and current local equivalents",[3.1,3.1])
h("10.2 Proposed team workflow")
p("Create the named repository on GitHub and connect the local repository to its remote. Invite the four group members through the repository's collaborator settings; each person accepts their invitation before contributing [4]. Agree on the requested branch names above before making further module-specific changes.")
p("Each member works on a focused branch, stages the relevant files and creates descriptive commits. After tests pass, the member pushes the branch and opens a pull request into main. Another member reviews the changes, discusses corrections and checks the test results. Merge after review and passing checks, retaining the pull request as the record of the integration decision [3].")
h("10.3 Commit process and checks")
p("Examples of suitable messages are ‘Implemented student linked list management’, ‘Added service request queue’, ‘Implemented AVL tree student sorting’, and ‘Created campus route graph BFS DFS’. These are examples for actual work, not claims that these four commits exist. The local baseline is a single implementation commit.")
p("The supplied PR template asks for change details and validation. The Java workflow compiles and tests with Java 17 on pushes and pull requests after publication. Local tests have passed, but remote CI execution and individual author history have not been verified. Appendix D provides separate evidence spaces for repository, branches, commits, pull requests and contributors.")

chapter(11,"Team Members and Contributions")
p("The following identities, responsibilities and contribution descriptions were supplied by the group for this report. They describe the stated allocation of work. Individual authorship is not independently established by the current single baseline commit; the group should attach matching commit/PR evidence in Appendix D before final submission.")
table(["Member Name","Student ID","Responsibility"],members,"Group members and assigned responsibilities",[1.65,1.55,3.0])
contributions=[
    ("J. Abdullah","23DA2-0575","Student Linked List",["Designed the Student class and its fields.","Implemented StudentNode and the custom StudentLinkedList.","Developed add, update, delete, search and display operations.","Tested student record management, including duplicate IDs and boundary deletions."],"Student.java; StudentNode.java; StudentLinkedList.java"),
    ("SMF. Asra","23DA2-0826","Queue and Stack",["Implemented the student service-request queue and enqueue/dequeue operations.","Implemented ActionStack and recent activity tracking.","Tested FIFO request order, LIFO action order, empty states and reuse."],"ServiceRequest.java; QueueNode.java; ServiceQueue.java; Action.java; StackNode.java; ActionStack.java"),
    ("MSM. Dilsath","23DA2-0576","Hash Table and BST/AVL",["Implemented the custom hash table and hashing-based student search.","Developed separate-chaining collision handling.","Developed BST/AVL insertion, searching and traversal.","Supported tree deletion, rebalancing and collision/resize tests."],"StudentHashTable.java; TreeNode.java; StudentBST.java; AVLTree.java"),
    ("MTF. Nifra","23DA2-0729","Campus Graph Routes",["Designed the campus graph structure with locations as vertices and roads as edges.","Developed add/remove location and route operations.","Implemented BFS and iterative DFS traversal.","Tested cyclic/disconnected routes and removal of incident roads."],"Location.java; Edge.java; CampusRouteGraph.java")
]
for i,(name,sid,role,items,classes) in enumerate(contributions,1):
    h("11."+str(i)+" "+name)
    p("Student ID: "+sid+". Responsibility: "+role+".")
    bullets(items)
    p("Assigned classes: "+classes+".")
h("11.5 Shared contribution")
p("The group reports shared participation in system integration, testing, debugging, documentation, GitHub collaboration and final demonstration preparation. The automated runner validates integrated behavior; the final collaboration evidence should identify the actual author, reviewer and commit or pull request for each contribution.")

chapter(12,"Testing and Results")
h("12.1 Test environment and method")
p("Testing was repeated against the current four-student version during report preparation on 24 September 2026. The Windows environment used JDK 23.0.2 and the supplied PowerShell script compiled sources with --release 17 and -Xlint:all. The dependency-free TestRunner completed 12 suites and 53,400 assertions without a reported failure. This is functional/invariant testing, not a performance benchmark or a remote CI run.")
p("The suite includes deterministic model-based tests: 600 record mutations are compared against an independent reference array; AVL tests inspect balance and heights during 400-key insertion and deletion sequences. Console tests feed scripted input through every main menu option and the additional submenus. Actual results below summarize these verified checks.")
table(["Test Case","Input","Expected Result","Actual Result","Status"],[
    ("Add student","Valid record; duplicate ID","Insert once in all indexes","Indexes agree; duplicate rejected","Pass"),
    ("Update student","Existing ID, changed name/GPA","All indexed references replaced","List/hash/tree match replacement","Pass"),
    ("Delete student","Existing ID, no waiting request","ID absent from all structures","No remaining indexed record","Pass"),
    ("Search","ID; repeated name; missing key","Correct matches or no result","Linear/binary/hash checks agree","Pass"),
    ("Queue","R1 then R2","R1 processed first","FIFO order and rear reset verified","Pass"),
    ("Stack","Add then update action","Update action pops first","LIFO and peek verified","Pass"),
    ("BST/AVL","Rotation cases; 400 keys","Sorted keys and balanced AVL","All invariants hold after mutations","Pass"),
    ("Hashing","AA, AR, RA, RR; 300 more keys","Collisions and resizing preserve data","Chain deletions and rehash verified","Pass"),
    ("Graph BFS","Cyclic A–E graph; isolated F","A, B, C, D, E; F alone","Expected BFS order observed","Pass"),
    ("Graph DFS","Same test graph","A, B, D, C, E","Expected DFS order observed","Pass"),
    ("Graph removal","Delete B and incident roads","No dangling references","Counts/traversals remain correct","Pass"),
    ("Invalid input","NaN, infinity, blank/menu errors","Reject and re-prompt safely","Scripted console checks pass","Pass"),
    ("Interrupted edit","EOF during add/update","No partial change","Indexes retain prior state","Pass"),
    ("Demo records","Launch --demo","Four records, 0001–0004","Four requested rows displayed","Pass")
],"Functional tests and observed results",[1.0,1.05,1.45,2.1,.6])
h("12.2 Sample records and observed ordering")
table(["Student ID","Name","Programme","GPA"],[("0001","Abdullah","BAIT","3.75"),("0002","Asra","IT","3.50"),("0003","Dilsath","Computer Science","3.90"),("0004","Nifra","Computer Science","3.60")],"Current sample student records",[1.1,1.25,2.95,.9])
p("AVL display returns 0001, 0002, 0003 and 0004. Descending GPA ranking returns Dilsath (3.90), Abdullah (3.75), Nifra (3.60) and Asra (3.50). The regenerated docs/sample-session.txt records actual console output for tree display, hash lookup, ranking, request processing, campus adjacency and both graph traversals.")
h("12.3 Test outcome and limitations")
p("The executed checks support the implemented functional requirements and data-structure invariants. They do not measure production-scale throughput, concurrent behavior, security or long-term persistence. The application is designed for an in-memory console demonstration. Target-runtime verification on Java 17 is configured in the GitHub workflow but has not been executed remotely in the inspected repository.")
placeholder("Automated test results","Insert terminal screenshot showing the 12 passing suites and 53,400 assertions.",1.0)

chapter(13,"Conclusion")
p("The University Student Record and Campus Route Management System integrates the required data structures into a working Java console application. Student records are maintained through a custom singly linked list, indexed by a custom hash table and organized by an AVL tree. Service requests demonstrate FIFO processing, while recent actions demonstrate LIFO history. Campus locations and roads are represented through a custom adjacency-list graph with BFS and DFS traversal.")
p("The objectives are supported by a complete 16-option menu, validated input, coordinated record changes, searching and sorting operations, and repeatable automated checks. The updated sample data and captured execution show how the modules work together. The implementation also illustrates trade-offs: hash lookup offers expected constant-time access, AVL traversal provides ID order, and linked-list mutations still require sequential work.")
p("The project provides a practical basis for explaining DSA concepts during a demonstration. Possible improvements include file or database persistence, authenticated roles, a graphical interface, request-status tracking and weighted shortest-path navigation using an additional algorithm. Such extensions should preserve the current separation between interface logic, coordinated operations and data-structure responsibilities.")
p("Before submission, the group should complete the lecturer and submission-date fields, insert the requested editor/terminal screenshots, confirm the lecture-material reference and attach genuine GitHub collaboration evidence. These additions concern submission evidence rather than missing core application functionality.")

chapter(14,"References")
p("The implementation analysis is based primarily on the supplied project source and executed tests. External references below support the Java platform context, standard DSA background and proposed GitHub workflow. No unprovided lecture material is represented as having been consulted.")
references=[
    "[1] Oracle. Java Platform, Standard Edition & Java Development Kit, Version 17 API Specification. Available: https://docs.oracle.com/en/java/javase/17/docs/api/ (accessed 24 September 2026).",
    "[2] M. T. Goodrich, R. Tamassia and M. H. Goldwasser. Data Structures and Algorithms in Java, 6th ed. Wiley, 2014. Publisher information: https://www.wiley-vch.de/en/areas-interest/computing-computer-sciences/data-structures-and-algorithms-in-java-978-1-118-80857-3. Background reading for data-structure theory; the report's implementation details are derived from project code.",
    "[3] GitHub. GitHub flow. Available: https://docs.github.com/en/get-started/using-github/github-flow (accessed 24 September 2026).",
    "[4] GitHub. Inviting collaborators to a personal repository. Available: https://docs.github.com/en/repositories/managing-your-repositorys-settings-and-features/repository-access-and-collaboration/inviting-collaborators-to-a-personal-repository (accessed 24 September 2026).",
    "[5] Sri Lanka Technological Campus. CIT300 - Data Structures and Algorithms lecture materials. [Complete lecturer, lecture titles, academic year and relevant pages from the materials actually used; these files were not supplied for report preparation.]",
    "[6] Project source and test artifacts. University Student Record and Campus Route Management System: src/, tests/TestRunner.java, README.md and docs/sample-session.txt. Local baseline commit 4128be8 with current sample-data edits, inspected and tested on 24 September 2026."
]
for reference in references: p(reference)

h("APPENDIX A: SOURCE CODE EVIDENCE",1)
p("The complete editable Java source is delivered in src/. Use the following placeholders for readable editor screenshots with the filename and relevant method visible. Each screenshot should support the implementation explanation rather than reproduce an entire unreadable source file.")
for title,detail in [
    ("Student model and validation","src/student/Student.java: fields, constructor and GPA validation."),
    ("Queue and stack operations","src/queue/ServiceQueue.java and src/stack/ActionStack.java: endpoint changes."),
    ("Hash collision handling","src/hashing/StudentHashTable.java: Entry chains, hash() and delete()."),
    ("AVL rotations","src/tree/AVLTree.java: restore(), rotateLeft() and rotateRight()."),
    ("Graph BFS and DFS code","src/graph/CampusRouteGraph.java: queue and stack traversal loops.")
]:
    placeholder(title,"Insert editor screenshot: "+detail,.9)

h("APPENDIX B: CAPTURED OUTPUT AND REPRODUCTION",1)
p("The excerpt below is text captured from a real demo run after the sample records were updated. It is editable output evidence, not a screenshot. The full transcript is docs/sample-session.txt, driven by docs/demo-input.txt. Additional screenshot spaces for all menu choices appear in Chapter 8.")
sample=(ROOT/"docs/sample-session.txt").read_text(encoding="utf-8-sig")
lines=sample.splitlines()
start=next(i for i,line in enumerate(lines) if line.startswith("Students using AVL inorder"))
end=next(i for i in range(start,len(lines)) if lines[i].startswith("Total students:"))+1
for line in lines[start:end]:
    para=p(line); para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.line_spacing=1.0; para.paragraph_format.space_after=Pt(2)
    for run in para.runs: run.font.name="Consolas"; run.font.size=Pt(9)
p("For the queue demonstration, student 0001 submits a transcript request. The recorded processing result is:")
p("Processed: REQ-1 | Student 0001 | Transcript | Printed transcript copy")
for line in lines:
    if "BFS traversal:" in line or "DFS traversal:" in line:
        p(line.replace("Starting location: ",""))
h("B.1 Reproduction commands")
bullets(["Build and test: powershell -NoProfile -ExecutionPolicy Bypass -File scripts/build.ps1 test", "Launch demo: java -cp out Main --demo", "Replay output: Get-Content docs/demo-input.txt | java -cp out Main --demo"])

h("APPENDIX C: UML AND DIAGRAM INDEX",1)
p("Editable diagram-generation code is supplied in reports/build_report.py. High-resolution diagram assets are in reports/assets/. The UML diagrams in Chapter 6 cover every required class, including the node classes supporting the list, queue, stack, trees and graph. The campus diagram in Chapter 7 uses the exact eight demo roads and distances from loadDemoData().")
table(["Diagram","Location in report / file"],[("System architecture","Chapter 3 / architecture.png"),("Student/list/hash UML","Section 6.1 / uml_students.png"),("BST/AVL UML","Section 6.1.1 / uml_trees.png"),("Queue/stack UML","Section 6.1.2 / uml_queue_stack.png"),("Graph UML","Section 6.1.3 / uml_graph.png"),("System flow","Section 6.2 / system_flow.png"),("Campus graph","Section 7.6 / campus_graph.png")],"Diagram source index",[2.25,3.95])
p("To revise wording, use the normal Word paragraphs and tables. To revise diagram structure, edit the corresponding drawing function in build_report.py and regenerate the document before refreshing the Word fields. Update the full table of contents, list of figures and list of tables after inserting final evidence, because page positions may change.")

h("APPENDIX D: GITHUB EVIDENCE",1)
p("These placeholders must be replaced by the group's actual GitHub evidence. No remote repository, pull-request or contributor screenshot has been invented. Keep branch names consistent with the final workflow and record any mapping from the current local names shown in Chapter 10.")
for title,detail in [
    ("GitHub repository","Repository title, owner/URL and source folders."),
    ("GitHub branches","main and the four agreed module branches."),
    ("GitHub commits","Meaningful commit messages, authors and dates."),
    ("GitHub pull requests","Module pull requests, reviewers and merge/check status."),
    ("GitHub contributors","Contributor identities and their recorded contributions.")
]:
    placeholder(title,"Insert actual screenshot: "+detail,1.0)

settings=document.settings.element
update=OxmlElement("w:updateFields"); update.set(qn("w:val"),"true"); settings.append(update)
document.save(OUT/"CIT300_Project_Report.docx")
print("Generated",OUT/"CIT300_Project_Report.docx")
print("Diagrams:",len(paths))
